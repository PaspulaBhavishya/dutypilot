import docx
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

SAMPLE_TIMETABLE_DATA = [
    ["satish@acse.edu", "Monday", "08:30", "09:30", "Advanced Database Systems"],
    ["satish@acse.edu", "Wednesday", "10:00", "11:00", "Big Data Analytics"],
    ["anitha@acse.edu", "Tuesday", "09:00", "10:30", "Software Engineering"],
    ["anitha@acse.edu", "Thursday", "14:00", "15:30", "Object Oriented Design"],
    ["rahul@acse.edu", "Monday", "10:30", "11:30", "Data Structures Lab"],
    ["rahul@acse.edu", "Friday", "08:15", "09:15", "Python Programming"],
    ["prasad@acse.edu", "Tuesday", "11:15", "12:15", "Operating Systems"],
    ["priya@acse.edu", "Monday", "14:00", "15:00", "Computer Networks"],
    ["madhav@ece.edu", "Monday", "09:00", "10:00", "Digital Signal Processing"],
    ["madhav@ece.edu", "Thursday", "11:00", "12:00", "VLSI Design"],
    ["sravani@ece.edu", "Monday", "10:30", "11:30", "Microprocessors"],
    ["krishna@ece.edu", "Wednesday", "09:00", "10:00", "Embedded Systems"],
    ["raghav@eee.edu", "Monday", "08:30", "09:30", "Power Electronics"],
    ["divya@eee.edu", "Monday", "11:00", "12:00", "Control Systems"],
    ["sekhar@mech.edu", "Tuesday", "10:00", "11:30", "Thermodynamics"],
    ["tarun@mech.edu", "Wednesday", "14:00", "15:00", "Fluid Mechanics"],
]

def generate_docx():
    doc = docx.Document()
    doc.add_heading("Faculty Daily Class Timetable", level=1)
    doc.add_paragraph("College Academic Timetable for Overlap Collision Detection & Exam Invigilation Scheduling.")
    
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    headers = ["Faculty Email", "Day of Week", "Start Time", "End Time", "Subject / Class Name"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        
    for row in SAMPLE_TIMETABLE_DATA:
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = val
            
    doc.save("sample_timetable.docx")
    print("Generated sample_timetable.docx")

def generate_pdf():
    pdf_doc = SimpleDocTemplate("sample_timetable.pdf", pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        spaceAfter=12,
        textColor=colors.HexColor('#1E293B')
    )
    elements.append(Paragraph("Faculty Daily Class Timetable", title_style))
    elements.append(Paragraph("College Academic Class Periods Schedule (Used for Minute-to-Minute Overlap Detection)", styles['Normal']))
    elements.append(Spacer(1, 14))
    
    table_data = [["Faculty Email", "Day of Week", "Start Time", "End Time", "Subject / Class Name"]]
    table_data.extend(SAMPLE_TIMETABLE_DATA)
    
    t = Table(table_data, colWidths=[150, 80, 70, 70, 150])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    elements.append(t)
    pdf_doc.build(elements)
    print("Generated sample_timetable.pdf")

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
