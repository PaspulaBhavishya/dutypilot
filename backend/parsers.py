import io
import re
from datetime import datetime
from typing import List, Dict, Any
import pandas as pd

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

DAYS_OF_WEEK = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

def normalize_time(t_str: str) -> str:
    """Normalizes arbitrary time representations to 24-hr 'HH:MM' string."""
    t_str = t_str.strip().upper()
    t_str = re.sub(r'\s+', ' ', t_str)
    
    formats = ["%I:%M %p", "%I:%M%p", "%I %p", "%H:%M", "%H.%M"]
    for fmt in formats:
        try:
            return datetime.strptime(t_str, fmt).strftime("%H:%M")
        except ValueError:
            pass
            
    match = re.search(r'(\d{1,2})[:\.](\d{2})', t_str)
    if match:
        h, m = int(match.group(1)), int(match.group(2))
        if "PM" in t_str and h < 12:
            h += 12
        elif "AM" in t_str and h == 12:
            h = 0
        return f"{h:02d}:{m:02d}"
        
    return "09:00"

def parse_time_range(range_str: str):
    """Splits time ranges like '09:00 - 10:00' or '9:30 AM to 11:00 AM'."""
    parts = re.split(r'[-–—]|to', range_str, flags=re.IGNORECASE)
    if len(parts) >= 2:
        return normalize_time(parts[0]), normalize_time(parts[1])
    return "09:00", "10:00"

def parse_docx_timetable(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Parses Word .docx file extracting class timetable rows."""
    if not docx:
        raise RuntimeError("python-docx is not installed.")
        
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted = []
    
    for table in doc.tables:
        if len(table.rows) < 2:
            continue
            
        header_cells = [c.text.strip().lower() for c in table.rows[0].cells]
        
        email_idx = next((i for i, h in enumerate(header_cells) if "email" in h or "faculty" in h or "name" in h), None)
        day_idx = next((i for i, h in enumerate(header_cells) if "day" in h), None)
        time_idx = next((i for i, h in enumerate(header_cells) if "time" in h or "slot" in h or "period" in h), None)
        start_idx = next((i for i, h in enumerate(header_cells) if "start" in h), None)
        end_idx = next((i for i, h in enumerate(header_cells) if "end" in h), None)
        subject_idx = next((i for i, h in enumerate(header_cells) if "subject" in h or "class" in h or "course" in h), None)
        
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if not cells:
                continue
                
            email = cells[email_idx] if email_idx is not None and email_idx < len(cells) else ""
            day = cells[day_idx] if day_idx is not None and day_idx < len(cells) else "Monday"
            subject = cells[subject_idx] if subject_idx is not None and subject_idx < len(cells) else "Lecture"
            
            matched_day = next((d for d in DAYS_OF_WEEK if d.lower() in day.lower()), "Monday")
            
            if start_idx is not None and end_idx is not None and start_idx < len(cells) and end_idx < len(cells):
                st = normalize_time(cells[start_idx])
                et = normalize_time(cells[end_idx])
            elif time_idx is not None and time_idx < len(cells):
                st, et = parse_time_range(cells[time_idx])
            else:
                st, et = "09:00", "10:00"
                
            if email:
                extracted.append({
                    "faculty_identifier": email,
                    "day_of_week": matched_day,
                    "class_name": subject or "Lecture",
                    "start_time": st,
                    "end_time": et,
                    "is_teaching": True
                })
                
    for p in doc.paragraphs:
        line = p.text.strip()
        parsed_entry = parse_text_line(line)
        if parsed_entry:
            extracted.append(parsed_entry)
            
    return extracted

def parse_pdf_timetable(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Parses PDF file by extracting text and matching tabular or line patterns."""
    if not PdfReader:
        raise RuntimeError("pypdf is not installed.")
        
    reader = PdfReader(io.BytesIO(file_bytes))
    extracted = []
    
    full_text = ""
    for page in reader.pages:
        txt = page.extract_text()
        if txt:
            full_text += txt + "\n"
            
    lines = [l.strip() for l in full_text.splitlines() if l.strip()]
    i = 0
    while i < len(lines):
        line = lines[i]
        
        # 1. Single line format (e.g. "email | Day | 09:00-10:00 | Subject")
        parsed_entry = parse_text_line(line)
        if parsed_entry:
            extracted.append(parsed_entry)
            i += 1
            continue
            
        # 2. Sequential cell format (common in PDF table extractions)
        email_match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', line)
        if email_match and i + 4 < len(lines):
            next_day_candidate = lines[i + 1]
            matched_day = next((d for d in DAYS_OF_WEEK if d.lower() == next_day_candidate.lower()), None)
            if matched_day:
                st = normalize_time(lines[i + 2])
                et = normalize_time(lines[i + 3])
                subject = lines[i + 4]
                extracted.append({
                    "faculty_identifier": email_match.group(0),
                    "day_of_week": matched_day,
                    "class_name": subject,
                    "start_time": st,
                    "end_time": et,
                    "is_teaching": True
                })
                i += 5
                continue
                
        i += 1
            
    return extracted

def parse_text_line(line: str) -> Dict[str, Any] | None:
    """Parses a single text line containing faculty timetable information."""
    if not line or len(line) < 10:
        return None
        
    matched_day = None
    for d in DAYS_OF_WEEK:
        if re.search(rf'\b{d}\b', line, re.IGNORECASE):
            matched_day = d
            break
            
    if not matched_day:
        return None
        
    time_match = re.search(r'(\d{1,2}[:\.]\d{2}(?:\s*[AaPp][Mm])?)\s*(?:[-–—]|to)\s*(\d{1,2}[:\.]\d{2}(?:\s*[AaPp][Mm])?)', line)
    if not time_match:
        return None
        
    start_t = normalize_time(time_match.group(1))
    end_t = normalize_time(time_match.group(2))
    
    email_match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', line)
    if email_match:
        faculty_id = email_match.group(0)
    else:
        parts = re.split(rf'\b{matched_day}\b', line, flags=re.IGNORECASE)
        faculty_id = parts[0].strip(' ,;|')
        
    remainder = line[time_match.end():].strip(' ,;|')
    subject = remainder if remainder else "Lecture"
    
    if faculty_id:
        return {
            "faculty_identifier": faculty_id,
            "day_of_week": matched_day,
            "class_name": subject,
            "start_time": start_t,
            "end_time": end_t,
            "is_teaching": True
        }
        
    return None

def parse_excel_or_csv_timetable(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Parses Excel or CSV files."""
    if filename.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(file_bytes))
    else:
        df = pd.read_excel(io.BytesIO(file_bytes))
        
    extracted = []
    cols = {c.strip().lower(): c for c in df.columns}
    
    email_col = next((cols[k] for k in cols if "email" in k or "faculty" in k or "name" in k), None)
    day_col = next((cols[k] for k in cols if "day" in k), None)
    start_col = next((cols[k] for k in cols if "start" in k), None)
    end_col = next((cols[k] for k in cols if "end" in k), None)
    time_col = next((cols[k] for k in cols if "time" in k or "slot" in k), None)
    subject_col = next((cols[k] for k in cols if "subject" in k or "class" in k or "course" in k), None)
    teaching_col = next((cols[k] for k in cols if "teaching" in k), None)
    
    for _, row in df.iterrows():
        email = str(row[email_col]).strip() if email_col and pd.notna(row[email_col]) else ""
        day = str(row[day_col]).strip() if day_col and pd.notna(row[day_col]) else "Monday"
        matched_day = next((d for d in DAYS_OF_WEEK if d.lower() in day.lower()), "Monday")
        subject = str(row[subject_col]).strip() if subject_col and pd.notna(row[subject_col]) else "Lecture"
        
        if start_col and end_col and pd.notna(row[start_col]) and pd.notna(row[end_col]):
            st = normalize_time(str(row[start_col]))
            et = normalize_time(str(row[end_col]))
        elif time_col and pd.notna(row[time_col]):
            st, et = parse_time_range(str(row[time_col]))
        else:
            st, et = "09:00", "10:00"
            
        is_t = True
        if teaching_col and pd.notna(row[teaching_col]):
            val = str(row[teaching_col]).strip().lower()
            is_t = val in ("1", "1.0", "true", "yes", "teaching")
            
        if email:
            extracted.append({
                "faculty_identifier": email,
                "day_of_week": matched_day,
                "class_name": subject,
                "start_time": st,
                "end_time": et,
                "is_teaching": is_t
            })
            
    return extracted
